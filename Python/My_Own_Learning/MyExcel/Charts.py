import openpyxl as xl
from openpyxl.chart import Reference, BarChart, pie_chart

wb = xl.load_workbook('PriceList.xlsx')
sheet = wb['Sheet1']

print('Input the values!')
for row in range(2, 7):
    val = int(input(f'Type the value for row {row}: '))
    sheet.cell(row, 4).value = val

print('Inputs successfully taken')

values = Reference(sheet, 4, 2 , 4, 6)
chart = BarChart()

chart.add_data(values)
sheet.add_chart(chart, 'e2')

wb.save('ZNewPriceList.xlsx')
print('File Successfully Saved')